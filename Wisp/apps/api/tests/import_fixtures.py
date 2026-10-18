"""Small fictional documents shared by extraction and browser tests."""
import io


def office(format):
    buffer = io.BytesIO()
    if format == 'docx':
        from docx import Document
        doc = Document()
        doc.add_paragraph('星光读书会😀')
        table = doc.add_table(rows=2, cols=2)
        for cell, text in zip([cell for row in table.rows for cell in row.cells], ['时间', '地点', '周六', '书店']):
            cell.text = text
        doc.add_paragraph('结束后整理书架。')
        doc.save(buffer)
    elif format == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '星光读书会😀'
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = '周六'; table.cell(0, 1).text = '书店'
        slide2 = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide2.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '结束后整理书架。'
        deck.save(buffer)
    elif format == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active; sheet.title = '活动安排'
        sheet.append(['事项', None, '人数'])
        sheet.append(['星光读书会😀', None, 12])
        sheet['C3'] = '=SUM(C2:C2)'
        hidden = book.create_sheet('备用'); hidden.sheet_state = 'hidden'; hidden['A1'] = '隐藏的材料'
        book.save(buffer); book.close()
    return buffer.getvalue()


def pdf(*, blank=False, encrypted=False, pages=1):
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=400, height=500)
        if not blank:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            content = DecodedStreamObject(); content.set_data(b'BT /F1 12 Tf 30 450 Td (Starlight book club) Tj ET')
            page[NameObject('/Contents')] = writer._add_object(content)
    if encrypted:
        writer.encrypt('test-password')
    buffer = io.BytesIO(); writer.write(buffer); writer.close()
    return buffer.getvalue()


if __name__ == '__main__':
    import sys
    from pathlib import Path
    target = Path(sys.argv[1]); target.mkdir(parents=True, exist_ok=True)
    for format in ('docx', 'pptx', 'xlsx'):
        (target / ('book-club.' + format)).write_bytes(office(format))
    (target / 'book-club.pdf').write_bytes(pdf())
    (target / 'blank.pdf').write_bytes(pdf(blank=True))
