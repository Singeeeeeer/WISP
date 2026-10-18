import asyncio
import io
import unittest
import zipfile
from unittest.mock import AsyncMock, patch
import httpx
from fastapi.testclient import TestClient
from app.config import Settings
from app.import_text import ImportFailure, Text, extract_text, parse_in_process
from import_fixtures import office, pdf

with patch('app.config.load_settings', return_value=Settings()):
    from app.main import create_app


class ExtractionTests(unittest.TestCase):
    def failure(self, code, data, format):
        with self.assertRaises(ImportFailure) as caught:
            extract_text(data, format)
        self.assertEqual(caught.exception.code, code)

    def test_docx_interleaved_paragraphs_table_unicode(self):
        value = extract_text(office('docx'), 'docx')
        self.assertEqual(value['schemaVersion'], 'import-text-1')
        self.assertEqual(value['text'], '星光读书会😀\n时间\t地点\n周六\t书店\n结束后整理书架。')
        self.assertTrue(value['warnings'])

    def test_nested_merged_docx_table_is_not_duplicated(self):
        from docx import Document
        d = Document(); table = d.add_table(rows=1, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1)); cell.text = 'merged'
        cell.add_table(rows=1, cols=1).cell(0, 0).text = 'nested'
        b = io.BytesIO(); d.save(b)
        text = extract_text(b.getvalue(), 'docx')['text']
        self.assertEqual(text.count('merged'), 1); self.assertEqual(text.count('nested'), 1)

    def test_pptx_slide_order_tables_and_groups(self):
        text = extract_text(office('pptx'), 'pptx')['text']
        self.assertEqual(text, '[幻灯片 1]\n星光读书会😀\n周六\t书店\n[幻灯片 2]\n结束后整理书架。')

    def test_xlsx_positions_formulas_hidden_sheet(self):
        value = extract_text(office('xlsx'), 'xlsx')
        self.assertIn('行号\tA\tB\tC', value['text'])
        self.assertIn('2\t星光读书会😀\t\t12', value['text'])
        self.assertIn('[公式，未计算] =SUM(C2:C2)', value['text'])
        self.assertIn('[工作表：备用（隐藏）]', value['text'])
        self.assertIn('隐藏', '\n'.join(value['warnings']))

    def test_xlsx_stale_dimensions_do_not_truncate_cells(self):
        import re
        original = office('xlsx')
        changed = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(changed, 'w', zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                content = source.read(item)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content, replacements = re.subn(rb'<dimension ref="[^"]+"\s*/>', b'<dimension ref="A1:A1"/>', content)
                    self.assertEqual(replacements, 1)
                target.writestr(item, content)
        self.assertEqual(extract_text(changed.getvalue(), 'xlsx')['text'], extract_text(original, 'xlsx')['text'])

    def test_xlsx_without_dimensions_labels_actual_columns(self):
        from openpyxl import Workbook
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet('streamed')
        sheet.append(['A heading', 'B heading', 'C heading'])
        sheet.append(['one', None, 'three'])
        content = io.BytesIO(); workbook.save(content)
        text = extract_text(content.getvalue(), 'xlsx')['text']
        self.assertIn('行号\tA\tB\tC', text)
        self.assertIn('2\tone\t\tthree', text)

    def test_pdf_pages_and_mixed_empty_pages(self):
        from pypdf import PdfReader, PdfWriter
        value = extract_text(pdf(pages=2), 'pdf')
        self.assertIn('[第 2 页]\nStarlight book club', value['text'])
        writer = PdfWriter(); writer.append(PdfReader(io.BytesIO(pdf())))
        writer.add_blank_page(width=400, height=500)
        b = io.BytesIO(); writer.write(b)
        mixed = extract_text(b.getvalue(), 'pdf')
        self.assertIn('第 2 页未提取到文字', '\n'.join(mixed['warnings']))

    def test_pdf_empty_encrypted_and_page_limit(self):
        self.failure('IMPORT_NO_TEXT', pdf(blank=True), 'pdf')
        self.failure('IMPORT_ENCRYPTED', pdf(encrypted=True), 'pdf')
        self.failure('IMPORT_LIMIT_EXCEEDED', pdf(blank=True, pages=301), 'pdf')

    def test_office_empty_documents_do_not_succeed_with_markers(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook
        for format, doc in [('docx', Document()), ('pptx', Presentation()), ('xlsx', Workbook())]:
            b = io.BytesIO(); doc.save(b)
            self.failure('IMPORT_NO_TEXT', b.getvalue(), format)

    def test_bad_format_corruption_and_byte_limit(self):
        self.failure('IMPORT_FORMAT_UNSUPPORTED', b'abc', 'doc')
        for format in ('docx', 'pptx', 'xlsx', 'pdf'):
            self.failure('IMPORT_FILE_INVALID', b'broken-file', format)
            self.failure('IMPORT_FILE_INVALID', b'', format)
        self.failure('IMPORT_FILE_INVALID', office('docx'), 'xlsx')
        self.failure('IMPORT_FILE_TOO_LARGE', b'x' * 20_000_001, 'pdf')
        self.failure('IMPORT_ENCRYPTED', b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + 'EncryptedPackage'.encode('utf-16-le'), 'docx')

    def test_xml_entities_macros_and_zip_budgets(self):
        original = office('docx')
        def add(name, content):
            b = io.BytesIO(original)
            with zipfile.ZipFile(b, 'a', compression=zipfile.ZIP_DEFLATED) as z:
                z.writestr(name, content)
            return b.getvalue()
        self.failure('IMPORT_FILE_INVALID', add('evil.xml', b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///secret">]><x>&x;</x>'), 'docx')
        self.failure('IMPORT_FORMAT_UNSUPPORTED', add('word/vbaProject.bin', b'macro'), 'docx')
        self.failure('IMPORT_LIMIT_EXCEEDED', add('huge.xml', b' ' * 12_000_001), 'docx')

    def test_utf16_nul_normalization_and_nonblank(self):
        text = Text(); text.add('😀' * 30000); self.assertEqual(len(text.finish()), 30000)
        with self.assertRaises(ImportFailure) as caught: text.add('x')
        self.assertEqual(caught.exception.code, 'IMPORT_TEXT_TOO_LARGE')
        text = Text(); text.add(' a\r\nb\rc\v '); self.assertEqual(text.finish(), ' a\nb\nc\n ')
        with self.assertRaises(ImportFailure): text.add('x\0')
        text = Text(); text.add('[第 1 页]', marker=True)
        with self.assertRaises(ImportFailure): text.finish()


class ImportApiTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(create_app(Settings(), transport=httpx.MockTransport(lambda _: self.fail('model called'))))

    def post(self, data, format='pdf'):
        return self.client.post('/api/import-text?format=' + format, content=data, headers={'Content-Type': 'application/octet-stream'})

    def test_real_child_four_formats_without_model_configuration(self):
        for format in ('docx', 'pptx', 'xlsx', 'pdf'):
            response = self.post(pdf() if format == 'pdf' else office(format), format)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['format'], format)
            self.assertIn('X-Request-ID', response.headers)
            self.assertLessEqual(len('\n'.join(response.json()['warnings'])), 2000)
        self.assertFalse(self.client.get('/api/health').json()['model_connected'])

    def test_errors_are_safe_and_model_byte_limit_unchanged(self):
        r = self.post(b'%PDF-invalid private-secret')
        self.assertEqual(r.status_code, 422)
        self.assertNotIn('private-secret', r.text)
        self.assertEqual(r.json()['error']['code'], 'IMPORT_FILE_INVALID')
        self.assertEqual(self.post(b'x', 'txt').status_code, 415)
        self.assertEqual(self.client.post('/api/import-text?format=pdf', json={}).status_code, 415)
        self.assertEqual(self.post(b'x' * 20_000_001).status_code, 413)
        self.assertEqual(self.client.post('/api/understand', content=b'x' * 400001).status_code, 413)

    def test_timeout_and_disconnect_kill_and_reap_child(self):
        async def run(disconnected):
            child = unittest.mock.Mock()
            child.returncode = None
            async def communicate(_):
                await asyncio.sleep(60)
            child.communicate = communicate
            child.wait = AsyncMock(return_value=0)
            child.kill = unittest.mock.Mock()
            async def receive():
                if not disconnected:
                    await asyncio.sleep(60)
                return {'type': 'http.disconnect'}
            request = unittest.mock.Mock(receive=receive)
            with patch('app.import_text.asyncio.create_subprocess_exec', AsyncMock(return_value=child)), patch('app.import_text.TIMEOUT', 30 if disconnected else 0):
                with self.assertRaises(asyncio.CancelledError if disconnected else ImportFailure):
                    await parse_in_process(b'x', 'pdf', request)
            child.kill.assert_called_once(); child.wait.assert_awaited_once()
        asyncio.run(run(False)); asyncio.run(run(True))

    def test_disconnect_through_http_middleware_stops_parser(self):
        async def run():
            app = create_app(Settings(), transport=httpx.MockTransport(lambda _: self.fail('model called')))
            child = unittest.mock.Mock(returncode=None)
            started = asyncio.Event()
            async def communicate(_):
                started.set()
                await asyncio.sleep(60)
            child.communicate = communicate
            child.wait = AsyncMock(return_value=0)
            child.kill = unittest.mock.Mock()
            incoming = asyncio.Queue()
            await incoming.put({'type': 'http.request', 'body': b'local-file', 'more_body': False})
            async def send(_):
                pass
            scope = {
                'type': 'http', 'asgi': {'version': '3.0'}, 'http_version': '1.1',
                'method': 'POST', 'scheme': 'http', 'path': '/api/import-text',
                'raw_path': b'/api/import-text', 'query_string': b'format=pdf', 'root_path': '',
                'headers': [(b'content-type', b'application/octet-stream')],
                'client': ('127.0.0.1', 1), 'server': ('127.0.0.1', 80),
            }
            with patch('app.import_text.asyncio.create_subprocess_exec', AsyncMock(return_value=child)):
                serving = asyncio.create_task(app(scope, incoming.get, send))
                try:
                    await asyncio.wait_for(started.wait(), 1)
                    await incoming.put({'type': 'http.disconnect'})
                    try:
                        await asyncio.wait_for(serving, 1)
                    except RuntimeError as error:
                        # BaseHTTPMiddleware may report the cancelled route's
                        # absent response after the client has already gone.
                        self.assertEqual(str(error), 'No response returned.')
                finally:
                    if not serving.done():
                        serving.cancel()
                    await asyncio.gather(serving, return_exceptions=True)
            child.kill.assert_called_once()
            child.wait.assert_awaited_once()
        asyncio.run(run())

    def test_cancelled_parse_cleans_child_and_disconnect_listener(self):
        async def run():
            child = unittest.mock.Mock(returncode=None)
            started = asyncio.Event()
            async def communicate(_):
                started.set()
                await asyncio.sleep(60)
            async def receive():
                await asyncio.sleep(60)
                return {'type': 'http.disconnect'}
            child.communicate = communicate
            child.wait = AsyncMock(return_value=0)
            child.kill = unittest.mock.Mock()
            before = asyncio.all_tasks()
            with patch('app.import_text.asyncio.create_subprocess_exec', AsyncMock(return_value=child)):
                parsing = asyncio.create_task(parse_in_process(b'local-file', 'pdf', unittest.mock.Mock(receive=receive)))
                await asyncio.wait_for(started.wait(), 1)
                parsing.cancel()
                with self.assertRaises(asyncio.CancelledError):
                    await parsing
            child.kill.assert_called_once()
            child.wait.assert_awaited_once()
            self.assertEqual(asyncio.all_tasks() - before, set())
        asyncio.run(run())

    def test_imported_sources_survive_all_three_gemini_contracts(self):
        from test_gemini import request, wire, output as check_generated
        from test_gemini_text import generated
        for format in ('docx', 'pptx', 'xlsx', 'pdf'):
            for op in ('check', 'understand', 'create'):
                body = request(); body['snapshot']['id'] = format + op
                body['snapshot']['source'] = format + '_file'
                body['snapshot']['origin'] = {'kind': format + '_file', 'addedAt': '2026-10-18T00:00:00Z', 'fileName': 'private.' + format, 'byteSize': 42, 'lastModified': 1, 'note': '解析提示'}
                if op == 'understand': body['summaryStyle'] = 'standard'
                def transport(r):
                    self.assertNotIn('private.' + format, r.content.decode())
                    return httpx.Response(200, json=wire(check_generated() if op == 'check' else generated(op)))
                client = TestClient(create_app(Settings(api_key='mock-key', model='mock-model'), transport=httpx.MockTransport(transport)))
                response = client.post('/api/' + op, json=body)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()['snapshot'], body['snapshot'])
